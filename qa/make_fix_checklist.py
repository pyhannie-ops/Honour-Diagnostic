from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL, OCHRE, BLACK = RGBColor(0x09, 0x56, 0x56), RGBColor(0xC0, 0x8A, 0x2E), RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
FONT = "Aptos"

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)
st = doc.styles["Normal"]
st.font.name, st.font.size, st.font.color.rgb = FONT, Pt(10), BLACK
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.15
st.paragraph_format.space_after = Pt(3)


def run(p, text, bold=False, color=None, size=None, italic=False):
    r = p.add_run(text)
    r.bold, r.italic, r.font.name = bold, italic, FONT
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return r


def heading(text, sub=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.keep_with_next = True
    run(p, text, bold=True, color=TEAL, size=13)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "C08A2E")):
        bt.set(qn("w:" + k), v)
    b.append(bt)
    pPr.append(b)
    if sub:
        q = doc.add_paragraph()
        q.paragraph_format.keep_with_next = True
        run(q, sub, italic=True, color=GREY, size=9)


def item(text, tag=None, why=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)
    p.paragraph_format.keep_together = True
    run(p, "☐  ", size=12, color=TEAL)
    if tag:
        run(p, f"[{tag}] ", bold=True, color=OCHRE, size=9)
    run(p, text)
    if why:
        run(p, "  " + why, italic=True, color=GREY, size=9)


def para(text, **kw):
    p = doc.add_paragraph()
    run(p, text, **kw)
    return p


t = doc.add_paragraph()
run(t, "Fix and Improve List: hannieverse.com", bold=True, color=TEAL, size=20)
para("Hannie Consultants cc  |  HFS Arc: Honour™, Flow™, SoftPower™  |  Honour™ Diagnostic Experience launches October 2026",
     color=OCHRE)
para("Tags: [V] = verified in my checks of the live site.  [T] = needs your test (I could not run these).  "
     "[R] = my recommendation.  Work top to bottom. Sections 1 to 3 come before launch. Sections 4 to 8 can follow in the first weeks.")
para("Date: ____________     Checked by: ____________________", size=9)

heading("1. Must do before launch", "Anything that could stop a visitor from booking, or that a sponsor will notice in the first minute.")
item("Test the Discovery Call booking end to end, on laptop and phone, through to the confirmation email", "T",
     "The single most important check. Confirm date, time zone and joining link are correct.")
item("Test the Honour™ \"Request to Book\" flow and confirm the visitor is told what happens next", "T")
item("Test the Facilitator Package booking flow end to end", "T")
item("Submit the homepage form and the Technical Writing form. Confirm the thank-you message shows and the email reaches you", "T")
item("Check every page on your phone: no sideways scrolling, no overlaps, no cut-off images, readable Honour™ table", "T")
item("Open the browser console (F12) on each page and note any red errors", "T")
item("Fix the About Us credential links. Several point to LinkedIn edit pages (\"add-edit/EDUCATION\") that only you can open", "V",
     "Visitors will hit a login wall. Use your public profile or a credential-verification link, or remove the links.")
item("Show the Honour™ fee, a range, or your founding-cohort price on the page, not only inside a PDF", "V",
     "Homepage, Honour and Book Online all say only \"Fee per participant\".")
item("Add one plain-language sentence near the top of the homepage saying what Honour™ is and that it launches in October", "R",
     "Example: \"A three-session, two-week diagnostic that shows where your team is misaligned, before you spend on training or change.\"")
item("Give the Facilitator Package card a duration, a one-line description and either a price or \"price on the call\"", "V",
     "It is the only service card on Book Online with no time listed.")
item("Confirm the three PDFs (Honour™ Brief, HFS Arc, fee sheet) are the final versions and open on a phone", "T")

heading("2. Bugs and polish I found", "Small things that make the site look unfinished. Quick to fix.")
item("Homepage hours read \"10:00am - 16:00 pm\". Use one format, for example 10:00 to 16:00 SAST", "V")
item("Honour page: \"I t is not a strategy session\" has a stray space (likely letter-spacing markup)", "V")
item("Use one company name pattern across header, body and footer (Hannie Consultants cc / HannieVerse Enterprise™)", "V")
item("Pick one voice: \"we\" on Home, Philosophy and About, but \"I\" on Honour licensing, Technical Writing and Connections", "V",
     "As a solo founder, \"I\" is honest. If you prefer \"we\", use it everywhere.")
item("Add alt text to the Our Philosophy compass image and any other pictures without it", "V",
     "Matters for your accessibility statement.")
item("Check the currency: the homepage shows \"USD ($)\" but the business is in South Africa. State which currency the fees are in", "V")
item("Click the three Connections \"Say Hello\" buttons. Two partner sites blocked my automated check and one did not respond", "T")
item("Give each Connections partner button a clearer label (for example \"Contact the editor\") instead of three identical \"Say Hello\"", "R")
item("Check the About Us \"2026\" pivot-year reference reads as intended", "T")
item("Check the Insights archive: it is not in the navigation or on any page I fetched. Add it, or remove it from your plans for launch", "V")
item("Book Online: standardise the button wording (\"Book Now\" vs \"Request to Book\") or explain the difference", "V")
item("Check the page titles and descriptions for every page (they show in Google results)", "T")

heading("3. Make the offer clearer to a sceptical buyer", "The biggest gap is proof and specifics. You are launching a new method, so be honest about that and replace proof with transparency.")
item("Say plainly that Honour™ is new and launches in October, and offer founding-cohort places (for example the first three engagements)", "R",
     "Honesty builds more trust than vague confidence. Never imply a track record you do not have.")
item("Add a sample deliverable to the Honour™ Brief: an \"illustrative\" findings page for a dummy team", "R",
     "This answers the buyer's main question: what do I get in my hand at the end?")
item("Describe each of the three sessions in one short paragraph: what happens, who attends, what the sponsor receives afterwards", "R",
     "This replaces most of the R·A·D and 8-step jargon with something concrete.")
item("Add a \"Is this for you? Is it not?\" block: the type of team, size, situation, and when to wait", "R",
     "Being clear about who it is not for makes the buyers it is for feel understood.")
item("Add a short FAQ: What if my team is remote? What if people will not be honest? Who sees the results? What happens after? Confidentiality?", "R")
item("Move your credentials (20+ years in financial-services operations and IT, SAFe®) higher, and connect them to the method", "R",
     "Right now your background is the strongest proof you have.")
item("Consider a small risk reducer for the founding cohort, such as a clear scope and a satisfaction promise (check it fits your refund policy)", "R")
item("Run 2 or 3 pilot engagements at a founding rate in exchange for feedback and permission for a testimonial", "R",
     "Even one short named quote is worth more than pages of copy.")
item("Add a plain-language line under the \"W T F? Willing To Flip?\" section, or move it lower on the page", "R",
     "The wordplay is memorable, but a corporate sponsor may read it as flippant before they know what you sell.")
item("Cut the poetic lines from the top of the homepage or move them below the offer", "R",
     "\"One step steady, one breath clear\" is charming after the visitor trusts you, but not before.")
item("Give Our Philosophy page a clear next step (a button to the Honour™ page or the discovery call)", "V",
     "It currently ends with no call to action.")
item("Add one concrete example on the Philosophy page of how a principle shows up in a session", "R",
     "\"We pace by breath, not urgency\" needs a real example to feel credible to a sponsor.")

heading("4. Present the whole HFS Arc as your business model", "The site serves three engagements, but only Honour™ exists at launch. Show the journey without making the site look unfinished.")
item("Show the three arcs as one journey on the homepage: Honour™ (see clearly), Flow™ (practise), SoftPower™ (lead)", "R",
     "One simple visual with three steps, Honour™ marked \"Open now\".")
item("Label Flow™ and SoftPower™ as \"Next in the Arc\" or a roadmap, not \"Coming Soon\" with no detail", "R",
     "A roadmap reads as a plan. \"Coming soon\" with no dates reads as unfinished.")
item("Add a one-paragraph promise for Flow™ and for SoftPower™: who it is for, what changes, roughly how long", "R",
     "Even a draft outline lets a sponsor picture the whole path.")
item("Add a \"Notify me\" or waitlist form for Flow™ and SoftPower™ so interest is captured instead of lost", "R",
     "Every Honour™ visitor who says yes to updates is a warm lead for the next arc.")
item("State clearly that Honour™ stands alone: a client can stop after it, or continue", "R",
     "Your HFS Arc page says each phase is a standalone engagement. Say it on the Honour™ page too.")
item("Explain how Honour™ results lead into Flow™: what the sponsor will know at the end and what choices it opens", "R",
     "This is your natural upsell and it helps the Honour™ price feel like the start of something, not a one-off.")
item("Keep the Arc order of release honest: Honour™ October 2026, then Flow™, then SoftPower™. Add target quarters only when you are confident", "R",
     "Missing a public date does more damage than having none.")
item("Decide whether the site name and navigation should be built around the Arc (Honour™ / Flow™ / SoftPower™ pages) or around services", "R",
     "When Flow™ launches, each arc will need its own page, brief and booking option.")
item("Plan the Book Online page so each new arc can be added as a service card without a redesign", "R")

heading("5. Calls to action and booking journey", "Make the next step obvious and low-risk on every page.")
item("Use one primary button on every page: \"Book a 30-minute call. No commitment, no pitch.\" It is your best line, so repeat it", "R")
item("On the Honour™ page, make Book a Discovery Call the main button and Download the Brief the secondary one", "R",
     "Two equal buttons compete for attention above the fold.")
item("Add a short \"What happens on the discovery call\" note next to the booking button (who, how long, what to prepare)", "R")
item("Add a \"Book\" button to the main navigation so it is visible from every page", "R")
item("Make the confirmation email warm and practical: joining link, what to prepare, how to reschedule, your contact", "R")
item("Add a reminder email or message the day before the call", "R")
item("Ask one or two useful questions at booking (team size, main challenge) so the call is focused", "R")
item("Make the Honour™ Brief download optional-email: ask for a work email in exchange, so you can follow up", "R")

heading("6. Facilitator licensing", "A second audience with a different need. Give them their own clear path.")
item("Add a short licensing overview: who it is for, what is included, how long it takes to get started", "R")
item("Show the licence model or a price range (per-engagement, as your page says), or say \"discussed on the call\"", "R")
item("Add a simple eligibility note (experience, OD or HR background) so the right people book", "R")
item("Add one line on how you protect the method: materials, trademarks, quality standards", "R",
     "Serious facilitators will ask before they commit.")
item("Consider a separate licensing page linked from the Honour™ page and the navigation", "R")

heading("7. Trust, legal and accessibility")
item("Read all eight policy pages once as a client would: Booking, Payment, Refund, Terms, Privacy, Disclaimer, Copyright, Accessibility", "T",
     "Make sure they match the launch offer, especially refunds, cancellations and data use for the diagnostic.")
item("Check privacy wording for the diagnostic itself: who sees participant responses, where they are stored, how long you keep them", "R",
     "Confidentiality is the first concern of any sponsor buying a team diagnostic.")
item("Add a visible privacy line at the point of the form and the download", "R")
item("Test the site with keyboard only (Tab key) and check text contrast against your accessibility statement", "T")
item("Add your company registration details in the footer if you want a corporate buyer to trust the business quickly", "R")
item("Make sure the LinkedIn links in the footer and About Us go to your public company page and profile", "V")

heading("8. Find-ability and measuring what works")
item("Set up basic analytics (page views, button clicks, form submissions, bookings) so you can see what leads to calls", "R")
item("Write a clear title and description for each page, using phrases a buyer would search (team alignment diagnostic, organisational alignment assessment)", "R")
item("Add your business to Google Business Profile if you want local search visibility in Cape Town", "R")
item("Publish 2 or 3 short articles on the Insights page about the problem Honour™ solves (misalignment, meeting narratives, why teams stall)", "R",
     "This is how new buyers find you before they know your name, and it shows how you think.")
item("Share the launch on LinkedIn with a short story of why you built Honour™", "R",
     "Your network is your first source of pilot clients.")
item("Record a 60 to 90 second video of you explaining Honour™ and place it near the top of the Honour™ page", "R",
     "For a solo consultant, seeing you builds trust faster than copy does.")

heading("9. After launch: first 90 days")
item("Deliver the first pilot engagements and ask for a short quote and permission to use it", "R")
item("Write a one-page anonymised case summary from the first engagement (with client permission)", "R")
item("Review analytics monthly: which pages lead to booked calls, and where visitors leave", "R")
item("Re-run the site checks (bookings, forms, phone view) after every significant change to the site", "T")
item("Build the Flow™ page and brief using the same pattern as Honour™: promise, session outline, sample deliverable, price, booking", "R")
item("Repeat the same approach for SoftPower™ when Flow™ is live", "R")

heading("Notes")
lines = [""] * 4
for _ in lines:
    p = doc.add_paragraph()
    run(p, "_" * 92, color=RGBColor(0xBB, 0xBB, 0xBB))


from docx.enum.text import WD_BREAK
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
t = doc.add_paragraph()
run(t, "Draft copy: What you get in each session", bold=True, color=TEAL, size=16)
para("Draft for your review. Text in [square brackets] is a placeholder. The mapping of the 8 steps to sessions is an assumption to correct.", italic=True, color=GREY, size=9)

def block(title, body, get=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    run(p, title, bold=True, color=OCHRE, size=11)
    q = doc.add_paragraph()
    run(q, body)
    if get:
        r = doc.add_paragraph()
        run(r, "What you get: ", bold=True, color=TEAL)
        run(r, get)

block("How the two weeks work",
      "Honour\u2122 runs over two weeks: three 90-minute online sessions and two short reflections you complete in your own time. Up to nine people take part. It does not fix or prescribe. It shows you what is happening so you can decide what to do next.")
block("Session 1: Reveal (Enter, Orient, Reveal)",
      "The team gets into the room and agrees how the work will run. Then each person is asked, in a structured way, what is happening in the team beyond the meeting narratives: how work really gets done, where it stalls and where people see things differently.",
      "A first honest picture of the team, in its own words, and a shared starting point that nobody has had to argue for. [Add anything handed to the sponsor after this session.]")
block("Reflection 1: Reflect (between Sessions 1 and 2, in your own time)",
      "Each participant privately reviews what surfaced, and adds what they did not say out loud.",
      "The things people held back in the room, so the next session works with the whole picture and not only the loudest voices. [Confirm format and time.]")
block("Session 2: Align",
      "The team looks at what came out and finds where they already agree and where they do not. Differences are named instead of smoothed over.",
      "A clear view of where the team is aligned and where it is misaligned, with the reasons. [Add any artefact produced here.]")
block("Reflection 2: Integrate (between Sessions 2 and 3, in your own time)",
      "Each person considers what the alignment picture means for their own role and what they would want to happen next.",
      "Individual commitments and questions, ready to bring to the final session.")
block("Session 3: Direct (Direct, Close)",
      "The team turns what it has seen into direction: what to keep, what to look at, and what could be next. It closes with a clear view of the choices open to you. Honour\u2122 does not prescribe solutions, so the direction belongs to the team.",
      "A shared set of next steps the team has chosen, and the sponsor's [summary or readout, format to be confirmed]. If further support helps, Flow\u2122 is the natural next step.")

heading("Questions to answer so this can become final copy")
for q in ["What does the sponsor receive at the end: a written report, a readout meeting, a one-page summary, or nothing beyond the sessions?",
          "Do the sponsor and participants see the same results, or does the sponsor see a summary only?",
          "Is the sponsor in the sessions, or do they stay out so people speak openly?",
          "What do you do between sessions (analysis, preparing the next session)?",
          "How long does each reflection take?"]:
    item(q)

doc.save("/home/user/Honour-Diagnostic/qa/Website-Fix-and-Improve-List.docx")
print("saved")
