from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL, OCHRE, BLACK = RGBColor(0x09, 0x56, 0x56), RGBColor(0xC0, 0x8A, 0x2E), RGBColor(0x1A, 0x1A, 0x1A)
FONT = "Aptos"

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)

st = doc.styles["Normal"]
st.font.name, st.font.size, st.font.color.rgb = FONT, Pt(10), BLACK
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.15
st.paragraph_format.space_after = Pt(3)


def run(p, text, bold=False, color=None, size=None):
    r = p.add_run(text)
    r.bold, r.font.name = bold, FONT
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return r


def heading(text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.keep_with_next = True
    run(p, text, bold=True, color=TEAL, size=size)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "C08A2E")):
        bt.set(qn("w:" + k), v)
    b.append(bt)
    pPr.append(b)


def item(text, note=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)
    p.paragraph_format.keep_together = True
    run(p, "☐  ", size=12, color=TEAL)
    run(p, text)
    if note:
        run(p, "  " + note, color=OCHRE)
        p.runs[-1].italic = True


def lines(n=2):
    for _ in range(n):
        p = doc.add_paragraph()
        run(p, "_" * 92, color=RGBColor(0xBB, 0xBB, 0xBB))


t = doc.add_paragraph()
run(t, "Website Check: hannieverse.com", bold=True, color=TEAL, size=20)
p = doc.add_paragraph()
run(p, "Hannie Consultants cc  |  Honour™ Diagnostic  |  Before the 1 October launch", color=OCHRE, size=10)
p = doc.add_paragraph()
run(p, "About 20 minutes. You need a laptop with Chrome and your phone. Tick each box as you go and write down anything odd next to it. "
       "Take a screenshot of every problem and send it back with this sheet.")
p = doc.add_paragraph()
run(p, "Date: ____________     Checked by: ____________________     Device / browser: ____________________", size=9)

heading("1. Laptop: every page and button")
for x in ["Launchpad (home): all buttons and links go where they say",
          "Our Philosophy: page loads, images show, nothing looks stretched",
          "Honour: \"Download: The Honour™ Brief\" opens the PDF",
          "Honour: \"Book a Discovery Call\" and \"Book a call\" (licensing) open the right booking page",
          "Honour: \"Fee per participant\" opens a PDF with the fee in it",
          "HFS Arc: \"Download: HFS Arc\" opens the PDF; \"View the Honour™ Diagnostic\" works",
          "About Us: click every credential link. Note any that show a LinkedIn login or an error",
          "Technical Writing: page loads and the enquiry form is visible",
          "Connections: click all three \"Say Hello\" buttons (Novelinks, LifePsych, Build to Strength)",
          "Footer on any page: open all 8 policy pages (Accessibility, Booking, Copyright, Disclaimer, Payment, Privacy, Refund, Terms)"]:
    item(x)

heading("2. Laptop: errors behind the scenes")
item("Press F12 (Mac: Cmd + Option + I), click \"Console\", then reload each page")
item("Screenshot any red messages, and write the page name next to them")
item("Look for words like \"Lorem ipsum\", \"placeholder\", or empty gaps where a picture should be")

heading("3. Forms: the money check")
item("Homepage \"have a question?\" form: fill it in with test details and click Submit", "Use the name \"QA Test\".")
item("Did the visitor see a thank-you message?")
item("Did the email arrive in your inbox? How long did it take?  ______________")
item("Technical Writing enquiry form: same test, same two questions")
item("Try submitting with one required field left empty. Does it show a clear error?")

heading("4. Booking flows: most important")
item("Book Online: Discovery Call. Pick a time, complete the booking, reach the confirmation")
item("Confirmation email arrived, with the right time, time zone, and a link to join")
item("Book Online: Honour™ \"Request to Book\". Is it clear what happens next?")
item("Book Online: Honour™ Facilitator Package. Does the card show a duration and a price?")
item("Cancel or reschedule from the email. Does the link work?")
item("Repeat the Discovery Call booking once on your phone")

heading("5. Phone: does anything break?")
for x in ["Open the homepage, Honour, HFS Arc, and Book Online on your phone",
          "Text never runs off the edge of the screen",
          "Try sliding the page sideways. It should not move",
          "No overlapping words, buttons, or pictures",
          "Pictures are not cut off or stretched",
          "Menu opens and every nav link works",
          "Buttons are easy to tap with a thumb",
          "The Honour™ table (fee, duration, sessions) is readable"]:
    item(x)

heading("6. Wording spot-check")
item("Homepage hours read \"10:00am - 16:00 pm\". Fix to one time format")
item("Honour page: \"I t is not a strategy session\" has a stray space")
item("Company name is the same everywhere (Hannie Consultants cc / HannieVerse Enterprise™)")
item("Any other typo or oddity you spot")

heading("Problems found")
p = doc.add_paragraph()
run(p, "Page  |  What happened  |  Screenshot taken?  |  Blocks a booking? (yes / no)", bold=True, size=9)
lines(6)

doc.save("/home/user/Honour-Diagnostic/qa/Website-Check-Hannieverse.docx")
print("saved")
